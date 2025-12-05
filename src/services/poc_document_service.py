import os
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


class POCDocumentService:
    def __init__(self, output_dir="data/generated_proposals"):
        self.output_dir = output_dir
        os.makedirs(self.output_dir, exist_ok=True)

    def create_poc_document(self, poc_content: dict) -> str:
        document = Document()

        # Título
        title = document.add_heading(poc_content.get("titulo", "Proposta de POC"), 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        document.add_paragraph(f"Data: {datetime.now().strftime('%d/%m/%Y')}")
        document.add_paragraph("Gerado por: noesys.ai Agent")

        # Resumo Executivo
        document.add_heading("Resumo Executivo", level=1)
        document.add_paragraph(poc_content.get("resumo_executivo", ""))

        # Escopo Técnico
        document.add_heading("Escopo Técnico", level=1)
        for item in poc_content.get("escopo_tecnico", []):
            document.add_paragraph(item, style="List Bullet")

        # Arquitetura Sugerida
        document.add_heading("Arquitetura Sugerida", level=1)
        document.add_paragraph(poc_content.get("arquitetura_sugerida", ""))

        # KPIs de Sucesso
        document.add_heading("KPIs de Sucesso", level=1)
        for kpi in poc_content.get("kpis_sucesso", []):
            document.add_paragraph(kpi, style="List Bullet")

        # Cronograma
        document.add_heading("Cronograma Estimado", level=1)
        document.add_paragraph(poc_content.get("cronograma_estimado", ""))

        # Salvar
        filename = f"POC_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        filepath = os.path.join(self.output_dir, filename)
        document.save(filepath)

        return filepath
