import logging
import os
from datetime import datetime

from docx import Document

from src.services.github_service import post_comment, upload_file_to_repo

# Configura logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


async def generate_and_attach_rcm_document(
    issue_number: int, analysis_data: dict, file_template_path: str
) -> str:
    """
    Preenche o template DOCX do RCM, salva o arquivo e o anexa à issue do GitHub.

    Args:
        issue_number: Número da issue do GitHub.
        analysis_data: Dados da análise (Escopo, PF, Estimativa, etc.) para preencher o RCM.
        file_template_path: Caminho para o template do RCM.

    Returns:
        O caminho para o arquivo DOCX gerado (local).
    """
    logger.info(f"Iniciando geração de RCM para issue #{issue_number}...")

    if not os.path.exists(file_template_path):
        raise FileNotFoundError(f"Template não encontrado: {file_template_path}")

    doc = Document(file_template_path)

    # Dados para preenchimento
    replacements = {
        "<Nº RCM>": str(issue_number),
        "<DATA>": datetime.now().strftime("%d/%m/%Y"),
        "<Escopo>": analysis_data.get("solucao_sugerida", "N/A"),
        "<Estimativa de Esforço>": analysis_data.get("nivel_esforco", "N/A"),
        "<Prazo>": analysis_data.get("esforco_resolucao_dias", "N/A"),
        "<PF>": str(
            analysis_data.get("detalhes_evolutiva", {}).get(
                "estimativa_pontos_funcao", "N/A"
            )
            if analysis_data.get("detalhes_evolutiva")
            else "N/A"
        ),
    }

    # Substituição simples em parágrafos
    for paragraph in doc.paragraphs:
        for key, value in replacements.items():
            if key in paragraph.text:
                paragraph.text = paragraph.text.replace(key, str(value))

    # Salvar o arquivo preenchido
    output_dir = "data/generated_rcm"
    os.makedirs(output_dir, exist_ok=True)
    generated_filename = f"RCM-{issue_number}.docx"
    generated_file_path = os.path.join(output_dir, generated_filename)

    doc.save(generated_file_path)
    logger.info(f"RCM salvo localmente em: {generated_file_path}")

    # Upload para o GitHub
    target_repo_path = f"docs/rcm/{generated_filename}"
    commit_message = f"Add RCM for issue #{issue_number}"

    try:
        file_url = await upload_file_to_repo(
            generated_file_path, target_repo_path, commit_message
        )
        logger.info(f"RCM upload realizado: {file_url}")

        # Comentar na issue
        comment_body = (
            f"### 📄 RCM Gerado\n\n"
            f"O Relatório de Controle de Mudança (RCM) foi gerado e anexado.\n"
            f"**Link para download/visualização:** [{generated_filename}]({file_url})\n\n"
            f"Por favor, revise e aprove para prosseguirmos."
        )
        await post_comment(issue_number, comment_body)

    except Exception as e:
        logger.error(f"Erro ao fazer upload ou comentar no GitHub: {e}")
        # Não falha o processo todo se o upload falhar, mas loga o erro
        # O arquivo local ainda existe

    return generated_file_path


async def generate_memoria_calculo(issue_number: int, analysis_data: dict) -> str:
    """
    Gera uma Memória de Cálculo (XLSX) simples com base nos dados da análise.
    """
    logger.info(
        f"Iniciando geração de Memória de Cálculo para issue #{issue_number}..."
    )

    try:
        import openpyxl
        from openpyxl import Workbook
    except ImportError:
        logger.error("openpyxl não instalado. Não é possível gerar XLSX.")
        return ""

    wb = Workbook()
    ws = wb.active
    ws.title = "Memória de Cálculo"

    # Cabeçalho
    ws["A1"] = "Memória de Cálculo - SISP"
    ws["A2"] = f"Issue: {issue_number}"
    ws["A3"] = f"Data: {datetime.now().strftime('%d/%m/%Y')}"

    # Dados
    ws["A5"] = "Item"
    ws["B5"] = "Valor"

    ws["A6"] = "Estimativa de Esforço"
    ws["B6"] = analysis_data.get("nivel_esforco", "N/A")

    ws["A7"] = "Pontos de Função (PF)"
    ws["B7"] = (
        analysis_data.get("detalhes_evolutiva", {}).get(
            "estimativa_pontos_funcao", "N/A"
        )
        if analysis_data.get("detalhes_evolutiva")
        else "N/A"
    )

    ws["A8"] = "Prazo (Dias)"
    ws["B8"] = analysis_data.get("esforco_resolucao_dias", "N/A")

    # Salvar
    output_dir = "data/generated_memcalc"
    os.makedirs(output_dir, exist_ok=True)
    filename = f"MemCalc-{issue_number}.xlsx"
    file_path = os.path.join(output_dir, filename)

    wb.save(file_path)
    logger.info(f"Memória de Cálculo salva localmente em: {file_path}")

    # Upload (Opcional, mas bom para consistência)
    target_repo_path = f"docs/memcalc/{filename}"
    try:
        file_url = await upload_file_to_repo(
            file_path, target_repo_path, f"Add MemCalc for issue #{issue_number}"
        )
        logger.info(f"MemCalc upload realizado: {file_url}")

        await post_comment(
            issue_number,
            f"### 📊 Memória de Cálculo Gerada\n\nArquivo: [{filename}]({file_url})",
        )
    except Exception as e:
        logger.error(f"Erro no upload da MemCalc: {e}")

    return file_path
