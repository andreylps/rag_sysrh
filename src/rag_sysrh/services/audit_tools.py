import logging
import re
from pathlib import Path
from typing import List, Tuple

import openpyxl
from docx import Document
from langchain_core.prompts import ChatPromptTemplate
from langchain_openai import ChatOpenAI

from rag_sysrh.core.knowledge_base import get_contexto_organizacional_completo
from rag_sysrh.engine.models import NonConformity, ResultadoSISP

logger = logging.getLogger(__name__)


class AuditTools:
    """
    Ferramentas de auditoria para o Agente de Qualidade & Compliance.
    """

    def __init__(self):
        self.llm = ChatOpenAI(model="gpt-4o", temperature=0)

    def audit_sisp_process_compliance(
        self, resultado: ResultadoSISP
    ) -> List[NonConformity]:
        """
        Verifica se os valores no Excel batem com o objeto de resultado (Motor Python).
        """
        issues = []

        # 1. Verifica se arquivos foram gerados
        if (
            not resultado.memoria_calculo_path
            or not Path(resultado.memoria_calculo_path).exists()
        ):
            issues.append(
                NonConformity(
                    type="PROCESS_SISP",
                    severity="CRITICAL",
                    description="Memória de Cálculo (.xlsx) não encontrada ou não gerada.",
                )
            )
            return issues

        # 2. Validação Numérica (Excel vs Objeto)
        try:
            wb = openpyxl.load_workbook(resultado.memoria_calculo_path, data_only=True)
            # Busca aba "CONTAGEM" (Case insensitive)
            sheet_name = next(
                (s for s in wb.sheetnames if s.upper() == "CONTAGEM"), None
            )

            if sheet_name:
                ws = wb[sheet_name]
                # Procura célula com valor do PF Líquido (busca heurística ou fixa)
                # Assumindo que o template tem uma estrutura onde o valor está explícito
                # Para este MVP, vamos verificar se o valor está presente em alguma célula
                found = False
                pf_liquido_str = str(resultado.pf_liquido_total)

                for row in ws.iter_rows():
                    for cell in row:
                        if cell.value is not None and str(cell.value) == pf_liquido_str:
                            found = True
                            break
                    if found:
                        break

                if not found:
                    # Tenta buscar por aproximação (float)
                    for row in ws.iter_rows():
                        for cell in row:
                            try:
                                val = float(cell.value)
                                if abs(val - resultado.pf_liquido_total) < 0.01:
                                    found = True
                                    break
                            except Exception:
                                pass
                        if found:
                            break

                if not found:
                    issues.append(
                        NonConformity(
                            type="PROCESS_SISP",
                            severity="HIGH",
                            description=f"Inconsistência Numérica: Valor do PF Líquido ({resultado.pf_liquido_total}) não encontrado na aba CONTAGEM do Excel.",
                        )
                    )
            else:
                issues.append(
                    NonConformity(
                        type="PROCESS_SISP",
                        severity="HIGH",
                        description="Aba 'CONTAGEM' não encontrada na Memória de Cálculo.",
                    )
                )

        except Exception as e:
            issues.append(
                NonConformity(
                    type="PROCESS_SISP",
                    severity="HIGH",
                    description=f"Erro ao ler Memória de Cálculo para validação: {e}",
                )
            )

        return issues

    def scan_for_security_risks(self, resultado: ResultadoSISP) -> List[NonConformity]:
        """
        Varre os arquivos gerados em busca de PII e termos inseguros.
        """
        issues = []
        files_to_scan = []
        if resultado.rcm_path and Path(resultado.rcm_path).exists():
            files_to_scan.append((resultado.rcm_path, "docx"))

        # Regex patterns
        patterns = {
            "CPF": r"\b\d{3}\.\d{3}\.\d{3}-\d{2}\b",
            "Email_Corporativo": r"\b[A-Za-z0-9._%+-]+@sc\.gov\.br\b",  # Exemplo
            "Senha_Explicita": r"(?i)senha\s*[:=]\s*\S+",
            "Termo_Inseguro": r"(?i)(desabilitar firewall|bypass|admin/admin|senha padrão)",
        }

        for file_path, file_type in files_to_scan:
            content = ""
            try:
                if file_type == "docx":
                    doc = Document(file_path)
                    content = "\n".join([p.text for p in doc.paragraphs])
                    for table in doc.tables:
                        for row in table.rows:
                            for cell in row.cells:
                                content += " " + cell.text

                # Scan
                for risk_name, pattern in patterns.items():
                    matches = re.findall(pattern, content)
                    if matches:
                        # Ofusca o match para o relatório
                        sample = matches[0][:3] + "***"
                        issues.append(
                            NonConformity(
                                type="SECURITY_RISK",
                                severity="CRITICAL"
                                if risk_name in ["CPF", "Senha_Explicita"]
                                else "MEDIUM",
                                description=f"Risco de Segurança ({risk_name}) detectado em {Path(file_path).name}: '{sample}'",
                            )
                        )

            except Exception as e:
                logger.error(f"Erro ao escanear segurança em {file_path}: {e}")

        return issues

    def evaluate_technical_clarity_llm(self, diagnostico: str) -> Tuple[float, str]:
        """
        Avalia a clareza técnica do diagnóstico usando LLM.
        Retorna (nota 0-5, crítica).
        """
        contexto_org = get_contexto_organizacional_completo()

        prompt = ChatPromptTemplate.from_messages(
            [
                (
                    "system",
                    f"""Você é um Auditor Técnico Sênior. Avalie a clareza e qualidade técnica do texto a seguir.
            
            {contexto_org}
            
            Critérios:
            1. Clareza: O texto é compreensível para um desenvolvedor?
            2. Objetividade: Vai direto ao ponto?
            3. Justificativa: Explica o 'porquê' das decisões técnicas?
            4. Contexto: O texto respeita as definições dos times e processos descritos acima?
            
            Responda no formato: NOTA (0.0 a 5.0) | CRÍTICA RESUMIDA
            """,
                ),
                ("user", "Texto para avaliação:\n{diagnostico}"),
            ]
        )

        try:
            response = self.llm.invoke(prompt.format(diagnostico=diagnostico))
            content = response.content.strip()

            # Parser simples
            parts = content.split("|")
            if len(parts) >= 2:
                try:
                    score = float(parts[0].strip())
                    critique = parts[1].strip()
                    return score, critique
                except Exception:
                    pass

            return 3.0, f"Avaliação automática inconclusiva. Resposta bruta: {content}"

        except Exception as e:
            return 0.0, f"Erro na avaliação de clareza: {e}"

    def verify_template_integrity(
        self, resultado: ResultadoSISP
    ) -> List[NonConformity]:
        """Verifica se os arquivos existem e têm tamanho > 0."""
        issues = []

        paths = [resultado.memoria_calculo_path, resultado.rcm_path]
        for p in paths:
            if p:
                path_obj = Path(p)
                if not path_obj.exists():
                    issues.append(
                        NonConformity(
                            type="TEMPLATE_INTEGRITY",
                            severity="CRITICAL",
                            description=f"Arquivo obrigatório não encontrado: {path_obj.name}",
                        )
                    )
                elif path_obj.stat().st_size == 0:
                    issues.append(
                        NonConformity(
                            type="TEMPLATE_INTEGRITY",
                            severity="CRITICAL",
                            description=f"Arquivo gerado está vazio: {path_obj.name}",
                        )
                    )
        return issues
