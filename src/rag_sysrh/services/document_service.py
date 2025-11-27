import logging
import os
from pathlib import Path
from typing import List, Optional

from pydantic import BaseModel

from rag_sysrh.engine.document_generator import DocumentGenerator
from rag_sysrh.engine.models import ItemFuncional, TipoFuncao
from rag_sysrh.engine.sisp_calculator import SISPCalculator

logger = logging.getLogger(__name__)


class DocumentGenerationRequest(BaseModel):
    titulo: str
    descricao: str
    itens: List[dict]  # Lista de itens funcionais simplificados
    deflator: float = 0.5


class DocumentGenerationResponse(BaseModel):
    rcm_filename: Optional[str] = None
    memoria_filename: Optional[str] = None
    message: str


class DocumentService:
    def __init__(self):
        self.calculator = SISPCalculator()
        # Assumindo que o diretório de templates está na raiz do projeto ou configurado
        # Ajuste o caminho conforme necessário. O DocumentGenerator usa "data/Layout" por padrão.
        self.generator = DocumentGenerator()

    def generate_documents(
        self, request: DocumentGenerationRequest
    ) -> DocumentGenerationResponse:
        try:
            # 1. Converter itens do request para ItemFuncional
            itens_funcionais = []
            for item in request.itens:
                itens_funcionais.append(
                    ItemFuncional(
                        nome=item.get("nome", "Item Sem Nome"),
                        tipo=TipoFuncao(item.get("tipo", "ALI")),
                        descricao=item.get("descricao", ""),
                        der_estimado=int(item.get("der", 0)),
                        rlr_estimado=int(item.get("rlr", 0)),
                        justificativa_contagem=item.get(
                            "justificativa", "Gerado via API"
                        ),
                    )
                )

            # 2. Calcular Métricas SISP
            resultado_sisp = self.calculator.calcular_pf(
                itens_funcionais, request.deflator
            )

            # 3. Gerar Documentos
            # Gera um ID temporário ou usa um sequencial
            # O DocumentGenerator busca o próximo ID do Neo4j, então podemos passar um dummy aqui
            # mas ele usa esse ID para criar a pasta.
            # Vamos deixar o gerador decidir o ID interno, mas precisamos passar algo.
            # O método gerar_rcm do generator já busca o ID sequencial.

            # Gerar RCM (Word)
            rcm_path = self.generator.gerar_rcm(
                solicitacao=request.titulo,
                diagnostico=request.descricao,
                resultado=resultado_sisp,
            )

            # Gerar Memória de Cálculo (Excel)
            # O gerador de memória precisa do ID. Vamos extrair do nome do arquivo RCM gerado ou gerar um novo.
            # O ideal seria refatorar o generator para ser mais coeso, mas vamos adaptar.
            # O gerar_rcm retorna o Path completo. Ex: data/resultado_analise/0001_2025/0001_2025 - Titulo.docx

            if str(rcm_path).startswith("ERRO") or str(rcm_path).startswith("TEMPLATE"):
                return DocumentGenerationResponse(
                    message=f"Erro na geração do RCM: {rcm_path}"
                )

            # Extrai o ID da pasta criada (parent)
            rcm_id_folder = rcm_path.parent.name  # Ex: 0001_2025

            memoria_path = self.generator.gerar_memoria_calculo(
                resultado_sisp, rcm_id_folder
            )

            return DocumentGenerationResponse(
                rcm_filename=str(rcm_path),  # Retorna o caminho relativo ou absoluto
                memoria_filename=str(memoria_path),
                message="Documentos gerados com sucesso.",
            )

        except Exception as e:
            logger.error(f"Erro no serviço de geração de documentos: {e}")
            return DocumentGenerationResponse(message=f"Erro interno: {str(e)}")

    def get_file_path(self, filename: str) -> Path:
        """
        Valida e retorna o caminho absoluto do arquivo solicitado para download.
        Segurança: Impede Path Traversal.
        """
        # O filename recebido pode ser um caminho relativo vindo do response anterior
        # Ex: data/resultado_analise/0001_2025/arquivo.docx

        base_path = Path(os.getcwd())
        file_path = base_path / filename

        # Resolve para caminho absoluto
        file_path = file_path.resolve()

        # Verifica se o arquivo está dentro do diretório do projeto (segurança básica)
        if not str(file_path).startswith(str(base_path)):
            raise ValueError("Acesso negado: Arquivo fora do diretório permitido.")

        if not file_path.exists():
            raise FileNotFoundError("Arquivo não encontrado.")

        return file_path

    # --- Mapeamento Global de Arquivos RCM (Em Memória) ---
    # Mapeia issue_number (int) -> file_path (str)
    RCM_FILE_MAP = {}

    def register_rcm_file(self, issue_number: int, file_path: str):
        """Registra o caminho do arquivo RCM para uma issue."""
        self.RCM_FILE_MAP[issue_number] = file_path

    def get_rcm_file(self, issue_number: int) -> Path:
        """Retorna o caminho do arquivo RCM associado a uma issue."""
        if issue_number not in self.RCM_FILE_MAP:
            raise FileNotFoundError(f"RCM não encontrada para a issue #{issue_number}")

        path_str = self.RCM_FILE_MAP[issue_number]
        return Path(path_str)

    def generate_simple_rcm_docx(self, issue_number: int, rcm_text: str) -> str:
        """
        Gera um DOCX simples a partir do texto do RCM e o registra.
        """
        from docx import Document

        doc = Document()
        doc.add_heading("Registro de Controle de Mudanças (RCM)", 0)

        # Adiciona o conteúdo
        # O texto vem em Markdown, idealmente faríamos um parse,
        # mas para simplicidade vamos adicionar como parágrafos.
        for line in rcm_text.split("\n"):
            if line.startswith("# "):
                doc.add_heading(line.replace("# ", ""), level=1)
            elif line.startswith("## "):
                doc.add_heading(line.replace("## ", ""), level=2)
            elif line.startswith("### "):
                doc.add_heading(line.replace("### ", ""), level=3)
            else:
                if line.strip():
                    doc.add_paragraph(line)

        # Salva o arquivo
        # Vamos salvar em data/rcms_fisicos para persistir um pouco mais que /tmp
        output_dir = Path("data/rcms_fisicos")
        output_dir.mkdir(parents=True, exist_ok=True)

        filename = f"RCM_{issue_number}.docx"
        file_path = output_dir / filename

        doc.save(file_path)

        # Registra no mapa global
        self.register_rcm_file(issue_number, str(file_path))

        return str(file_path)

    # --- Mapeamento Global de Arquivos Memória de Cálculo (Em Memória) ---
    MEMCALC_FILE_MAP = {}

    def register_memcalc_file(self, issue_number: int, file_path: str):
        """Registra o caminho do arquivo Memória de Cálculo para uma issue."""
        self.MEMCALC_FILE_MAP[issue_number] = file_path

    def get_memcalc_file(self, issue_number: int) -> Path:
        """Retorna o caminho do arquivo Memória de Cálculo associado a uma issue."""
        if issue_number not in self.MEMCALC_FILE_MAP:
            # Fallback: Tenta gerar on-the-fly se não existir (para facilitar testes)
            # Em produção real, deveria ser gerado no fluxo.
            # Vamos gerar um dummy aqui se não existir.
            return Path(self.generate_memory_of_calculation_xlsx(issue_number, {}))

        path_str = self.MEMCALC_FILE_MAP[issue_number]
        return Path(path_str)

    def generate_memory_of_calculation_xlsx(
        self, issue_number: int, data: dict = None
    ) -> str:
        """
        Gera uma Memória de Cálculo em XLSX e a registra.
        """
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        wb = Workbook()
        ws = wb.active
        ws.title = "Memória de Cálculo"

        # Estilos
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(
            start_color="4F81BD", end_color="4F81BD", fill_type="solid"
        )
        center_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # Cabeçalhos
        headers = [
            "Item",
            "Descrição",
            "Tipo",
            "Estimativa (Horas)",
            "Custo (PF)",
            "Justificativa",
        ]
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        # Dados de Exemplo (Placeholder ou vindos de 'data')
        # Se 'data' for fornecido, usaríamos aqui. Por enquanto, placeholders.
        rows = [
            (
                1,
                "Análise de Requisitos",
                "Documentação",
                4.0,
                0.5,
                "Levantamento inicial",
            ),
            (
                2,
                "Implementação Backend",
                "Codificação",
                16.0,
                2.0,
                "Desenvolvimento da API",
            ),
            (
                3,
                "Implementação Frontend",
                "Codificação",
                12.0,
                1.5,
                "Telas e Componentes",
            ),
            (4, "Testes Unitários", "Qualidade", 8.0, 1.0, "Cobertura de testes"),
            (
                5,
                "Deploy e Validação",
                "Infraestrutura",
                2.0,
                0.25,
                "Implantação em Homologação",
            ),
        ]

        for row_idx, row_data in enumerate(rows, 2):
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                if col_idx in [4, 5]:  # Numéricos
                    cell.alignment = center_align

        # Ajuste de largura das colunas
        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 40
        ws.column_dimensions["C"].width = 20
        ws.column_dimensions["D"].width = 20
        ws.column_dimensions["E"].width = 15
        ws.column_dimensions["F"].width = 40

        # Salva o arquivo
        output_dir = Path("data/memoria_calculo_fisica")
        output_dir.mkdir(parents=True, exist_ok=True)

        filename = f"Memoria_Calculo_{issue_number}.xlsx"
        file_path = output_dir / filename

        wb.save(file_path)

        # Registra no mapa global
        self.register_memcalc_file(issue_number, str(file_path))

        return str(file_path)
