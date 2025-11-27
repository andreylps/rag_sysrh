import os
import shutil

from docx import Document

# Configurações de Caminhos
# Em um cenário real, isso viria de variáveis de ambiente ou config.py
SIMULATOR_DOCS_PATH = "D:/Projeto IA/PROJETOS/rh-gov-simulador/docs"
LOCAL_DOCS_ARCHIVE_PATH = "data/manuais"


def generate_or_update_operational_manual(issue: dict, issue_type: str) -> str:
    """
    Gera ou atualiza o Manual Operacional relacionado à issue.
    Salva uma cópia no simulador e outra no arquivo local.
    """
    # 1. Determinar o nome do arquivo do manual
    # Por simplicidade, vamos criar um manual único por enquanto ou baseado no título?
    # O prompt sugere "Geração/Atualização", implicando que pode existir.
    # Vamos usar um nome genérico "Manual_Operacional_SYSRH.docx" ou algo derivado da issue?
    # Se for "Atualização", idealmente seria um manual geral.
    # Vamos assumir um manual geral do sistema por enquanto, ou um por módulo se tivéssemos essa info.
    # Vamos criar um manual específico para a feature da issue para evitar conflitos de merge complexos agora.
    # Ex: "Manual_Operacional_Feature_X.docx"

    sanitized_title = (
        "".join(
            c
            for c in issue.get("title", "Feature")
            if c.isalnum() or c in (" ", "_", "-")
        )
        .strip()
        .replace(" ", "_")
    )
    manual_filename = f"Manual_Operacional_{sanitized_title}.docx"

    # Garante que os diretórios existem
    os.makedirs(SIMULATOR_DOCS_PATH, exist_ok=True)
    os.makedirs(LOCAL_DOCS_ARCHIVE_PATH, exist_ok=True)

    simulator_file_path = os.path.join(SIMULATOR_DOCS_PATH, manual_filename)
    local_archive_file_path = os.path.join(LOCAL_DOCS_ARCHIVE_PATH, manual_filename)

    # 2. Gerar Conteúdo (Simulado aqui, mas viria da LLM)
    # Em um agente real, chamaríamos a LLM para escrever o manual com base na descrição da issue e diffs.
    content = f"""
    # Manual Operacional - {issue.get("title")}
    
    ## Visão Geral
    Esta funcionalidade foi implementada/ajustada na Issue #{issue.get("number")}.
    
    ## Descrição
    {issue.get("body", "Sem descrição fornecida.")}
    
    ## Instruções de Uso
    1. Acesse o sistema.
    2. Navegue até a funcionalidade.
    3. Execute a operação conforme descrito nos requisitos.
    
    ## Solução de Problemas
    Em caso de erro, verifique os logs do sistema ou contate o suporte.
    
    ---
    Gerado automaticamente pelo Agente de Documentação.
    """

    # 3. Criar ou Atualizar o Documento
    # Verifica se já existe no simulador (fonte da verdade)
    if os.path.exists(simulator_file_path):
        doc = Document(simulator_file_path)
        doc.add_page_break()
        doc.add_heading(f"Atualização - Issue #{issue.get('number')}", level=1)
        # Adiciona o novo conteúdo
        for line in content.split("\n"):
            if line.strip():
                doc.add_paragraph(line.strip())
    else:
        doc = Document()
        doc.add_heading(f"Manual Operacional - {issue.get('title')}", 0)
        for line in content.split("\n"):
            if line.strip():
                doc.add_paragraph(line.strip())

    # 4. Salvar em AMBOS os locais
    try:
        doc.save(simulator_file_path)
        # Copia para o local archive para garantir sincronia exata
        shutil.copy2(simulator_file_path, local_archive_file_path)

        # Retorna o caminho do simulador para uso posterior (validação)
        return simulator_file_path
    except Exception as e:
        raise RuntimeError(f"Erro ao salvar manual: {str(e)}")
