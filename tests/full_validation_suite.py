import asyncio
import logging
import os
import sys
from datetime import datetime
from unittest.mock import patch

# Adiciona o diretório raiz ao path para importações
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from src.agents.dev_agent import run_dev_agent
from src.services.github_service import (
    create_issue,
    get_issue_details,
    update_issue_labels,
)
from src.services.rcm_generation_service import (
    generate_and_attach_rcm_document,
    generate_memoria_calculo,
)

# Configuração de Logging
logging.basicConfig(
    level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s"
)
logger = logging.getLogger("FullValidationSuite")


async def test_scenario_1_happy_path():
    """
    CENÁRIO 1: O "Caminho Feliz" da Fábrica
    """
    logger.info("\n=== CENÁRIO 1: CAMINHO FELIZ (EVOLUTIVA + DOC APROVADA) ===")

    # 1. Criar Issue
    title = f"[TESTE #1] Implementar Relatório de Cursos {datetime.now().strftime('%H%M%S')}"
    body = "Como gestor, preciso de um relatório de cursos realizados pelos servidores."
    issue_number = await create_issue(title, body)
    logger.info(f"Issue criada: #{issue_number}")

    # 2. Agente Analista (Simulado/Rápido)
    # Vamos assumir que o analista rodou e gerou o relatório.
    # Para o teste, vamos chamar diretamente a geração de RCM para garantir o estado.
    analysis_data = {
        "solucao_sugerida": "Implementar relatório em PDF com lista de cursos.",
        "nivel_esforco": "Médio",
        "esforco_resolucao_dias": "5",
        "detalhes_evolutiva": {"estimativa_pontos_funcao": 10},
    }

    # Gerar RCM
    template_path = "data/templates/SysRH - RCM - Relatório de Controle de Mudança - MOD - 9999-9999.docx"
    # Criar template dummy se não existir
    if not os.path.exists(template_path):
        os.makedirs(os.path.dirname(template_path), exist_ok=True)
        from docx import Document

        doc = Document()
        doc.add_paragraph("RCM Template <Nº RCM>")
        doc.save(template_path)

    await generate_and_attach_rcm_document(issue_number, analysis_data, template_path)

    # Aprovar RCM
    await update_issue_labels(issue_number, add_labels=["status:pronto-para-dev"])
    logger.info("RCM gerado e aprovado. Issue pronta para dev.")

    # 3. Agente Desenvolvedor (Real)
    # O dev agent vai gerar um manual. Esperamos que seja bom.
    logger.info("Executando Agente Desenvolvedor...")
    result = await run_dev_agent(issue_number, {"title": title, "body": body})
    logger.info(f"Resultado Dev: {result}")

    # 4. Verificação QA
    details = await get_issue_details(issue_number)
    labels = details["labels"]

    if "status:aguardando-correcao-doc" not in labels:
        logger.info("✅ CENÁRIO 1 SUCESSO: QA Aprovou (ou não reprovou).")
    else:
        logger.error("❌ CENÁRIO 1 FALHA: QA Reprovou o manual do caminho feliz.")


async def test_scenario_2_unhappy_path():
    """
    CENÁRIO 2: O "Caminho Infeliz" (Falha Crítica na Documentação)
    """
    logger.info("\n=== CENÁRIO 2: CAMINHO INFELIZ (VIOLAÇÃO DE SEGURANÇA) ===")

    # 1. Criar Issue
    title = f"[TESTE #2] Corrigir Bug e Adicionar Detalhe de Login {datetime.now().strftime('%H%M%S')}"
    body = "Preciso corrigir o login e documentar a senha de root."
    issue_number = await create_issue(title, body)

    # Pular para Dev
    await update_issue_labels(issue_number, add_labels=["status:pronto-para-dev"])

    # 2. Agente Desenvolvedor (Com Sabotagem)
    # Vamos mockar a função que gera o manual para retornar um arquivo com a senha proibida.

    # Criar arquivo sabotado
    bad_manual_path = f"data/generated_docs/Manual_Sabotado_{issue_number}.docx"
    os.makedirs(os.path.dirname(bad_manual_path), exist_ok=True)
    from docx import Document

    doc = Document()
    doc.add_heading("Manual Operacional", 0)
    doc.add_paragraph(
        "A senha de acesso ao banco de dados de testes é 'root123'."
    )  # VIOLAÇÃO
    doc.save(bad_manual_path)

    logger.info(f"Arquivo sabotado criado em: {bad_manual_path}")

    # Patch no dev_agent para usar nosso arquivo ruim
    # O caminho do patch deve ser onde a função é IMPORTADA no dev_agent.py
    with patch(
        "src.agents.dev_agent.generate_or_update_operational_manual",
        return_value=bad_manual_path,
    ):
        logger.info("Executando Agente Desenvolvedor (Sabotado)...")
        result = await run_dev_agent(issue_number, {"title": title, "body": body})
        logger.info(f"Resultado Dev: {result}")

    # 3. Verificação QA
    details = await get_issue_details(issue_number)
    labels = details["labels"]

    if "status:aguardando-correcao-doc" in labels:
        logger.info("✅ CENÁRIO 2 SUCESSO: QA Reprovou o manual e bloqueou a issue.")
    else:
        logger.error("❌ CENÁRIO 2 FALHA: QA Aprovou um manual com senha exposta!")


async def test_scenario_3_artifacts():
    """
    CENÁRIO 3: Geração de Artefatos Físicos
    """
    logger.info("\n=== CENÁRIO 3: GERAÇÃO DE ARTEFATOS FÍSICOS ===")

    # Usando uma issue nova para garantir limpeza
    issue_number = await create_issue(
        f"[TESTE #3] Artefatos {datetime.now().strftime('%H%M%S')}",
        "Teste de artefatos",
    )

    analysis_data = {
        "solucao_sugerida": "Teste de Artefatos",
        "nivel_esforco": "Baixo",
        "esforco_resolucao_dias": "1",
        "detalhes_evolutiva": {"estimativa_pontos_funcao": 5},
    }

    # 1. Gerar RCM
    template_path = "data/templates/SysRH - RCM - Relatório de Controle de Mudança - MOD - 9999-9999.docx"
    rcm_path = await generate_and_attach_rcm_document(
        issue_number, analysis_data, template_path
    )

    # 2. Gerar Memória de Cálculo
    memcalc_path = await generate_memoria_calculo(issue_number, analysis_data)

    # Verificação Física
    if os.path.exists(rcm_path):
        logger.info(f"✅ RCM criado fisicamente: {rcm_path}")
    else:
        logger.error(f"❌ RCM não encontrado: {rcm_path}")

    if os.path.exists(memcalc_path):
        logger.info(f"✅ Memória de Cálculo criada fisicamente: {memcalc_path}")
    else:
        logger.error(f"❌ Memória de Cálculo não encontrada: {memcalc_path}")


async def main():
    logger.info("Iniciando Bateria de Testes Finais - Missão Auditoria...")
    try:
        await test_scenario_1_happy_path()
        await test_scenario_2_unhappy_path()
        await test_scenario_3_artifacts()
    except Exception as e:
        logger.error(f"Erro fatal na bateria de testes: {e}", exc_info=True)


if __name__ == "__main__":
    asyncio.run(main())
