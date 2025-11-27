import os

from docx import Document


def create_rcm_template():
    doc = Document()
    doc.add_heading("Relatório de Controle de Mudança (RCM)", 0)

    doc.add_paragraph("RCM Nº: <Nº RCM>")
    doc.add_paragraph("Data: <DATA>")

    doc.add_heading("1. Descrição da Mudança", level=1)
    doc.add_paragraph("<Escopo>")

    doc.add_heading("2. Estimativas", level=1)
    doc.add_paragraph("Estimativa de Esforço: <Estimativa de Esforço>")
    doc.add_paragraph("Prazo Estimado: <Prazo>")
    doc.add_paragraph("Pontos de Função: <PF>")

    doc.add_heading("3. Aprovação", level=1)
    doc.add_paragraph("Aprovação do Cliente: _____________________________")

    os.makedirs("data/templates", exist_ok=True)
    output_path = "data/templates/SysRH - RCM - Relatório de Controle de Mudança - MOD - 9999-9999.docx"
    doc.save(output_path)
    print(f"Template created at: {output_path}")


if __name__ == "__main__":
    create_rcm_template()
